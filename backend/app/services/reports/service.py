"""Report generation (PDF, CSV, Excel) from the same services that power the UI.

Reports are generated with the *requesting user's* permissions and store scope (re-checked at build
time), stored with a SHA-256 checksum, and carry the data date, sources and the synthetic-data notice.
"""

from __future__ import annotations

import csv
import hashlib
import io
import uuid
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import OrgContext
from app.core.errors import NotFoundError, PermissionDeniedError, ValidationAppError
from app.core.permissions import Permission, Role, permissions_for
from app.models import Membership, Notification, Organization, Report, User
from app.services.storage import report_path

P = Permission
REPORT_TYPES: dict[str, dict[str, Any]] = {
    "executive": {"title": "Executive summary", "permission": P.DASHBOARD_READ},
    "financial": {"title": "Financial analysis (P&L)", "permission": P.FINANCE_READ},
    "inventory": {"title": "Inventory & supplier report", "permission": P.INVENTORY_READ},
    "forecast": {"title": "Forecast performance report", "permission": P.FORECASTS_READ},
    "scenario_comparison": {"title": "Scenario comparison", "permission": P.SCENARIOS_READ},
    "benchmarking": {"title": "Company benchmarking", "permission": P.INTELLIGENCE_READ},
}
FORMATS = ("pdf", "csv", "xlsx")


def context_for(db: Session, org_id: uuid.UUID, user_id: uuid.UUID) -> OrgContext:
    row = db.execute(select(Membership, Organization, User).join(Organization, Organization.id == Membership.org_id)
                     .join(User, User.id == Membership.user_id)
                     .where(Membership.org_id == org_id, Membership.user_id == user_id)).first()
    if row is None or row[0].status != "active":
        raise PermissionDeniedError("The report requester is no longer an active member")
    membership, org, user = row
    role = Role(membership.role)
    return OrgContext(org=org, membership=membership, user=user, role=role, permissions=permissions_for(role),
                      store_scope=list(membership.store_scope) if membership.store_scope else None)


# ---------------------------------------------------------------------------------------------
# Content builders -> {"title", "subtitle", "sections": [{"heading", "paragraphs", "table": {"columns", "rows"}}]}
# ---------------------------------------------------------------------------------------------


def _money(v: Any) -> str:
    from app.services.advisor.formatting import money

    return money(v) if isinstance(v, int | float) else "—"


def _pct(v: Any) -> str:
    return f"{v * 100:.1f}%" if isinstance(v, int | float) else "—"


def _num(v: Any, d: int = 1) -> str:
    return f"{v:,.{d}f}" if isinstance(v, int | float) else "—"


def build_executive(db: Session, ctx: OrgContext, params: dict[str, Any]) -> dict[str, Any]:
    from app.services.analytics.dashboard import executive_dashboard

    d = executive_dashboard(db, ctx, _uuid(params.get("company_id")), _uuids(params.get("store_ids")),
                            params.get("period"), _date(params.get("start")), _date(params.get("end")))
    kpi_rows = []
    for k in d["kpis"]:
        fmt = _money if k["unit"] == "currency" else _pct if k["unit"] == "percent" else _num
        kpi_rows.append([k["label"], fmt(k["value"]), fmt(k["comparison"]),
                         _pct(k["change"]) if k.get("change_kind") == "relative" else (fmt(k["change"]) if k["change"] is not None else "—"),
                         (k.get("status") or "") + (" (estimated)" if k.get("estimated") else "")])
    sections = [{"heading": "Key performance indicators",
                 "table": {"columns": ["Metric", "Value", "Comparison", "Change", "Status"], "rows": kpi_rows}}]
    if d.get("health"):
        sections.append({"heading": f"Business health score: {d['health']['score']} ({d['health']['grade']})",
                         "table": {"columns": ["Component", "Raw value", "Score", "Weight"],
                                   "rows": [[c["key"], _num(c["raw"], 3), _num(c["score"]), _pct(c["effective_weight"])]
                                            for c in d["health"]["components"]]},
                         "paragraphs": d["health"]["limitations"]})
    sections.append({"heading": "Store performance",
                     "table": {"columns": ["Store", "Net sales", "Growth", "Contribution margin", "Sales/sq ft (yr)"],
                               "rows": [[f"{s['code']} {s['name']}", _money(s["net_sales"]), _pct(s.get("growth")),
                                         _pct(s.get("store_contribution_margin")), _num(s.get("sales_per_sqft_annualized"), 0)]
                                        for s in d["stores"]]}})
    sections.append({"heading": "Categories", "table": {"columns": ["Category", "Net sales", "Share", "Gross margin", "Growth"],
                                                        "rows": [[c["category"], _money(c["net_sales"]), _pct(c["share"]),
                                                                  _pct(c.get("gross_margin")), _pct(c.get("growth"))]
                                                                 for c in d["categories"]]}})
    sections.append({"heading": "Data freshness and sources",
                     "table": {"columns": ["Data", "Latest date", "Sources"],
                               "rows": [[f["fact"], f["latest_date"] or "no data",
                                         ", ".join(s["name"] for s in f.get("sources", []))] for f in d["freshness"]]},
                     "paragraphs": d["notes"]})
    return {"title": "Executive summary", "subtitle": f"{d['company']['name']} · {d['period']['label']}",
            "data_end": d["data_end"], "synthetic": d["company"]["is_synthetic"], "sections": sections,
            "primary_table": sections[0]["table"]}


def build_financial(db: Session, ctx: OrgContext, params: dict[str, Any]) -> dict[str, Any]:
    from app.services.analytics import operations as ops
    from app.services.analytics.dashboard import resolve_scope_period
    from app.services.analytics.pnl import compute_pnl, monthly_pnl_series

    scope, p, data_end = resolve_scope_period(db, ctx, _uuid(params.get("company_id")), _uuids(params.get("store_ids")),
                                              params.get("period"), _date(params.get("start")), _date(params.get("end")))
    months = monthly_pnl_series(db, scope, p.start, p.end, data_end)
    total = compute_pnl(db, scope, p.start, p.end, data_end)
    cols = ["Month", "Net sales", "COGS", "Gross profit", "Gross margin", "Operating expenses", "EBITDA", "Net profit",
            "Net margin"]
    rows = [[mo["label"] + (" (partial)" if mo["partial"] else ""), _money(mo["net_sales"]), _money(mo["cogs"]),
             _money(mo["gross_profit"]), _pct(mo["gross_margin"]), _money(mo["operating_expenses"]), _money(mo["ebitda"]),
             _money(mo["net_profit"]), _pct(mo["net_margin"])] for mo in months]
    rows.append(["Total", _money(total.net_sales), _money(total.cogs), _money(total.gross_profit), _pct(total.gross_margin),
                 _money(total.operating_expenses), _money(total.ebitda), _money(total.net_profit), _pct(total.net_margin)])
    bva = ops.budget_vs_actual(db, scope, p.start, p.end, data_end)
    sections = [{"heading": "Monthly profit & loss", "table": {"columns": cols, "rows": rows}, "paragraphs": total.notes},
                {"heading": "Expenses by category",
                 "table": {"columns": ["Category", "Amount"],
                           "rows": [[k, _money(v)] for k, v in sorted(total.expenses_by_category.items(), key=lambda kv: -kv[1])]}}]
    if bva.get("status") == "ok":
        sections.append({"heading": "Budget vs actual",
                         "table": {"columns": ["Month", "Metric", "Budget", "Actual", "Variance", "Favourable"],
                                   "rows": [[r["month"][:7], r["metric"], _money(r["budget"]), _money(r["actual"]),
                                             _pct(r["variance_pct"]), "yes" if r["favourable"] else "no"]
                                            for r in bva["rows"][:60]]}})
    else:
        sections.append({"heading": "Budget vs actual", "paragraphs": [bva.get("note") or "No budget data."]})
    return {"title": "Financial analysis", "subtitle": f"{scope.company.name} · {p.label}", "data_end": data_end.isoformat(),
            "synthetic": scope.company.is_synthetic, "sections": sections, "primary_table": sections[0]["table"]}


def build_inventory(db: Session, ctx: OrgContext, params: dict[str, Any]) -> dict[str, Any]:
    from app.services.analytics import operations as ops
    from app.services.analytics.dashboard import resolve_scope_period

    scope, p, data_end = resolve_scope_period(db, ctx, _uuid(params.get("company_id")), _uuids(params.get("store_ids")),
                                              params.get("period"), None, None)
    risks = ops.stockout_and_overstock(db, scope, data_end, limit=40)
    sup = ops.supplier_performance(db, scope, p.start, p.end) if ctx.has(P.SUPPLIERS_READ) else {"suppliers": []}
    risk_cols = ["Store", "SKU", "Product", "On hand", "On order", "Avg daily demand", "Days of cover", "Lead time", "Risk"]
    sections = [
        {"heading": f"Stockout risks (as of {risks.get('as_of')})", "paragraphs": [risks.get("method", "")],
         "table": {"columns": risk_cols, "rows": [[r["store_code"], r["sku"], r["product"], _num(r["on_hand_qty"], 0),
                                                   _num(r.get("on_order_qty"), 0), _num(r.get("avg_daily_demand")),
                                                   _num(r.get("days_of_cover")), _num(r.get("lead_time_days")), r["risk"]]
                                                  for r in risks.get("stockout_risks", [])]}},
        {"heading": f"Overstock (excess value {_money(risks.get('overstock_value_total'))})",
         "table": {"columns": ["Store", "SKU", "Product", "On hand", "Days of cover", "Excess units", "Excess value"],
                   "rows": [[r["store_code"], r["sku"], r["product"], _num(r["on_hand_qty"], 0), _num(r.get("days_of_cover")),
                             _num(r.get("excess_units"), 0), _money(r.get("excess_value"))] for r in risks.get("overstock", [])]}},
        {"heading": "Supplier lead times",
         "table": {"columns": ["Supplier", "Orders", "Avg lead (days)", "P90 lead", "Contracted", "On time", "Alert"],
                   "rows": [[s["supplier"], s["orders"], _num(s["avg_lead_days"]), _num(s["p90_lead_days"]),
                             _num(s["contracted_lead_days"], 0), _pct(s["on_time_rate"]), "yes" if s["alert"] else ""]
                            for s in sup.get("suppliers", [])]}},
    ]
    return {"title": "Inventory & supplier report", "subtitle": f"{scope.company.name} · {p.label}",
            "data_end": data_end.isoformat(), "synthetic": scope.company.is_synthetic, "sections": sections,
            "primary_table": sections[0]["table"]}


def build_forecast(db: Session, ctx: OrgContext, params: dict[str, Any]) -> dict[str, Any]:
    from app.models import ForecastModel
    from app.services.forecasting.service import model_points

    company = ctx.default_company(db, _uuid(params.get("company_id")))
    models = db.scalars(select(ForecastModel).where(ForecastModel.company_id == company.id,
                                                    ForecastModel.is_active.is_(True))).all()
    if ctx.store_scope is not None:
        models = [mm for mm in models if mm.level == "store" and uuid.UUID(mm.level_key) in ctx.store_scope]
    rows = [[mm.target, mm.level, mm.level_key[:8] if mm.level == "store" else mm.level_key, mm.algorithm, f"v{mm.version}",
             _pct(mm.metrics.get("wape")), _pct(mm.baseline_metrics.get("wape")), _pct(mm.skill_vs_baseline),
             _pct(mm.metrics.get("coverage_80")), mm.reliability, mm.training_end.isoformat()] for mm in models]
    sections = [{"heading": "Model registry (active versions)",
                 "table": {"columns": ["Target", "Level", "Key", "Algorithm", "Version", "WAPE", "Baseline WAPE", "Skill",
                                       "80% coverage", "Reliability", "Trained through"], "rows": rows},
                 "paragraphs": ["WAPE = Σ|actual − forecast| ÷ Σ|actual| on rolling-origin backtests (no future data). "
                                "Skill = 1 − WAPE ÷ best baseline WAPE."]}]
    company_model = next((mm for mm in models if mm.target == "net_sales" and mm.level == "company"), None)
    if company_model:
        pts = model_points(db, company_model)
        pts["month"] = pts["period_start"].map(lambda dd: dd.replace(day=1))
        monthly = pts.groupby("month").agg(yhat=("yhat", "sum"), lo=("lower_80", "sum"), hi=("upper_80", "sum")).reset_index()
        sections.append({"heading": "Company net sales forecast (monthly)",
                         "table": {"columns": ["Month", "Forecast", "Lower (80%, summed)", "Upper (80%, summed)"],
                                   "rows": [[r.month.isoformat()[:7], _money(r.yhat), _money(r.lo), _money(r.hi)]
                                            for r in monthly.head(12).itertuples(index=False)]},
                         "paragraphs": company_model.limitations})
    return {"title": "Forecast performance report", "subtitle": company.name, "data_end": None,
            "synthetic": company.is_synthetic, "sections": sections, "primary_table": sections[0]["table"]}


def build_scenarios(db: Session, ctx: OrgContext, params: dict[str, Any]) -> dict[str, Any]:
    from app.models import Scenario, ScenarioRun

    ids = _uuids(params.get("scenario_ids")) or []
    query = select(Scenario).where(Scenario.org_id == ctx.org_id, Scenario.status != "archived")
    if ids:
        query = query.where(Scenario.id.in_(ids))
    rows = []
    for s in db.scalars(query.order_by(Scenario.updated_at.desc()).limit(12)).all():
        run = db.scalar(select(ScenarioRun).where(ScenarioRun.scenario_id == s.id).order_by(ScenarioRun.computed_at.desc()))
        if run is None:
            continue
        r = run.result
        d = (r.get("delta") or {}).get("totals") or {}
        unc = (r.get("uncertainty") or {}).get("net_profit_delta") or {}
        inv = r.get("investment") or {}
        rows.append([s.name, ", ".join(lv["type"] for lv in s.levers), _money(d.get("net_sales")), _money(d.get("net_profit")),
                     _money(unc.get("p10")), _money(unc.get("p90")), _money(inv.get("investment")),
                     str(inv.get("payback_months") or "—"), str(len(r.get("warnings", [])))])
    sections = [{"heading": "Scenario comparison",
                 "table": {"columns": ["Scenario", "Levers", "Δ net sales", "Δ net profit", "P10", "P90", "Investment",
                                       "Payback (months)", "Warnings"], "rows": rows},
                 "paragraphs": ["Δ = scenario minus forecast baseline over each scenario's horizon. P10/P90 from Monte "
                                "Carlo simulation of estimated parameters. Ranking basis for risk-adjusted outcome: P10."]}]
    return {"title": "Scenario comparison", "subtitle": ctx.org.name, "data_end": None, "synthetic": ctx.org.is_demo,
            "sections": sections, "primary_table": sections[0]["table"]}


def build_benchmarking(db: Session, ctx: OrgContext, params: dict[str, Any]) -> dict[str, Any]:
    from app.services.intelligence import benchmark

    company = ctx.default_company(db) if ctx.has(P.FINANCE_READ) and ctx.store_scope is None else None
    b = benchmark(db, ctx.org_id, company, params.get("company_keys"))
    rows = [[r["display_name"] + (" (synthetic)" if r.get("is_synthetic") else ""), r.get("fiscal_year") or "—",
             r.get("currency") or "—", _pct(r.get("revenue_growth")), _pct(r.get("gross_margin")), _pct(r.get("ebitda_margin")),
             _pct(r.get("net_margin")), _num(r.get("store_count"), 0), _num(r.get("inventory_turnover")),
             _num((r.get("health") or {}).get("score"))] for r in b["rows"]]
    sections = [{"heading": "Benchmark", "table": {"columns": ["Company", "Fiscal year", "Currency", "Revenue growth",
                                                               "Gross margin", "EBITDA margin", "Net margin", "Stores",
                                                               "Inventory turnover", "Health score"], "rows": rows},
                 "paragraphs": b["comparability_notes"]}]
    return {"title": "Company benchmarking", "subtitle": ctx.org.name, "data_end": None,
            "synthetic": any(r.get("is_synthetic") for r in b["rows"]), "sections": sections,
            "primary_table": sections[0]["table"]}


BUILDERS = {"executive": build_executive, "financial": build_financial, "inventory": build_inventory,
            "forecast": build_forecast, "scenario_comparison": build_scenarios, "benchmarking": build_benchmarking}


# ---------------------------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------------------------


def render_pdf(content: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title=content["title"], author="AUREX")
    styles = getSampleStyleSheet()
    navy = colors.HexColor("#0F1F3D")
    h1 = ParagraphStyle("h1", parent=styles["Title"], textColor=navy, alignment=0, fontSize=20, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=navy, fontSize=13, spaceBefore=10)
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#33415C"))
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=7.5, leading=9)
    story: list[Any] = [Paragraph(content["title"], h1), Paragraph(_esc(content.get("subtitle", "")), styles["Heading4"])]
    meta = f"Generated {datetime.now(UTC):%Y-%m-%d %H:%M} UTC"
    if content.get("data_end"):
        meta += f" · data through {content['data_end']}"
    story.append(Paragraph(meta, small))
    if content.get("synthetic"):
        story.append(Paragraph("<b>SYNTHETIC DEMONSTRATION DATA</b> — fictional company NovaMart; not real business data.",
                               ParagraphStyle("warn", parent=small, textColor=colors.HexColor("#9A3412"))))
    for section in content["sections"]:
        story.append(Paragraph(_esc(section["heading"]), h2))
        table = section.get("table")
        if table and table["rows"]:
            data = [[Paragraph(f"<b>{_esc(c)}</b>", cell) for c in table["columns"]]]
            data += [[Paragraph(_esc(str(v)), cell) for v in row] for row in table["rows"]]
            t = Table(data, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF9")),
                                   ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C7D2E5")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FC")])]))
            story.append(t)
        elif table is not None:
            story.append(Paragraph("No rows for this section.", small))
        for paragraph in section.get("paragraphs") or []:
            if paragraph:
                story.append(Spacer(1, 2))
                story.append(Paragraph(_esc(paragraph), small))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Figures are computed by AUREX from stored data with documented formulas. Estimates and "
                           "forecasts carry the limitations stated in each section.", small))
    doc.build(story)
    return buffer.getvalue()


def _esc(text: str) -> str:
    return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render_csv(content: dict[str, Any]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([f"# {content['title']} - {content.get('subtitle', '')}"])
    if content.get("synthetic"):
        writer.writerow(["# SYNTHETIC DEMONSTRATION DATA - fictional company NovaMart"])
    for section in content["sections"]:
        table = section.get("table")
        if not table:
            continue
        writer.writerow([])
        writer.writerow([f"## {section['heading']}"])
        writer.writerow(table["columns"])
        writer.writerows(table["rows"])
    return buffer.getvalue().encode("utf-8-sig")


def render_xlsx(content: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "About"
    ws["A1"] = content["title"]
    ws["A1"].font = Font(bold=True, size=14, color="0F1F3D")
    ws["A2"] = content.get("subtitle", "")
    ws["A3"] = f"Generated {datetime.now(UTC):%Y-%m-%d %H:%M} UTC" + (f"; data through {content['data_end']}"
                                                                      if content.get("data_end") else "")
    if content.get("synthetic"):
        ws["A4"] = "SYNTHETIC DEMONSTRATION DATA - fictional company NovaMart; not real business data."
        ws["A4"].font = Font(bold=True, color="9A3412")
    used = set()
    for section in content["sections"]:
        table = section.get("table")
        if not table:
            continue
        name = "".join(ch for ch in section["heading"] if ch not in "[]:*?/\\")[:28] or "Sheet"
        while name in used:
            name = name[:25] + str(len(used))
        used.add(name)
        sheet = wb.create_sheet(name)
        sheet.append(table["columns"])
        for c in sheet[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8EEF9")
        for row in table["rows"]:
            sheet.append(list(row))
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(10, width + 2), 48)
        notes_row = sheet.max_row + 2
        for paragraph in section.get("paragraphs") or []:
            sheet.cell(row=notes_row, column=1, value=paragraph).alignment = Alignment(wrap_text=False)
            notes_row += 1
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


RENDERERS = {"pdf": render_pdf, "csv": render_csv, "xlsx": render_xlsx}


def build_report(db: Session, report_id: uuid.UUID) -> dict[str, Any]:
    report = db.get(Report, report_id)
    if report is None:
        raise NotFoundError("Report not found")
    report.status = "running"
    db.commit()
    try:
        ctx = context_for(db, report.org_id, report.created_by)
        spec = REPORT_TYPES[report.report_type]
        ctx.require(P.REPORTS_GENERATE, spec["permission"])
        content = BUILDERS[report.report_type](db, ctx, report.parameters or {})
        data = RENDERERS[report.format](content)
        path = report_path(report.org_id, report.id, report.format)
        Path(path).write_bytes(data)
        report.file_path = str(path)
        report.file_size = len(data)
        report.sha256 = hashlib.sha256(data).hexdigest()
        report.status = "completed"
        report.completed_at = datetime.now(UTC)
        report.error = None
        db.add(Notification(org_id=report.org_id, user_id=report.created_by, kind="report_ready",
                            title=f"Report ready: {report.title}", body=f"{report.format.upper()} · {len(data):,} bytes",
                            link=f"/app/reports?highlight={report.id}"))
        db.commit()
        return {"report_id": str(report.id), "status": "completed", "bytes": len(data)}
    except Exception as exc:
        db.rollback()
        report = db.get(Report, report_id)
        report.status = "failed"
        report.error = f"{type(exc).__name__}: {getattr(exc, 'message', str(exc))}"[:1000]
        db.commit()
        raise


def validate_request(ctx: OrgContext, report_type: str, fmt: str) -> dict[str, Any]:
    if report_type not in REPORT_TYPES:
        raise ValidationAppError("Unknown report type", details={"allowed": list(REPORT_TYPES)})
    if fmt not in FORMATS:
        raise ValidationAppError("Unknown format", details={"allowed": list(FORMATS)})
    ctx.require(P.REPORTS_GENERATE, REPORT_TYPES[report_type]["permission"])
    return REPORT_TYPES[report_type]


def _uuid(value: Any) -> uuid.UUID | None:
    return uuid.UUID(str(value)) if value else None


def _uuids(values: Any) -> list[uuid.UUID] | None:
    return [uuid.UUID(str(v)) for v in values] if values else None


def _date(value: Any) -> date | None:
    return date.fromisoformat(str(value)) if value else None
